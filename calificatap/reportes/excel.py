"""
excel.py — Exporta los resultados a un archivo de Excel con formato.

Hojas del archivo:
    Resumen        → datos del examen, promedios y gráfico por materia
    Estudiantes    → una fila por estudiante (verde aprueba / rojo no aprueba)
    Materias       → estadísticas de cada materia
    Preguntas      → % de acierto y opciones marcadas en cada pregunta
    Respuestas     → lo que marcó cada estudiante (verde bien / rojo mal)
    Alertas        → cosas para revisar (si las hay)

pandas escribe los datos y openpyxl les da el formato (colores, anchos, gráfico).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calificar import Resultados, pesos_saber

VERDE = "1F7A4D"
VERDE_SUAVE = "E3F1E9"
ROJO_SUAVE = "FBE3E1"
AMARILLO_SUAVE = "FFF4D6"
GRIS = "F2F4F3"


def _formatear_hoja(hoja, anchos: dict[str, int] | None = None, fila_titulos: int = 1) -> None:
    """Títulos en verde, columnas a la medida y primera fila fija."""
    for celda in hoja[fila_titulos]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[fila_titulos].height = 32
    hoja.freeze_panes = hoja.cell(row=fila_titulos + 1, column=1)
    for i, columna in enumerate(hoja.iter_cols(min_row=fila_titulos), start=1):
        letra = get_column_letter(i)
        largo = max((len(str(c.value)) for c in columna if c.value is not None), default=6)
        hoja.column_dimensions[letra].width = (anchos or {}).get(letra, min(max(largo + 2, 7), 45))


def exportar_excel(r: Resultados, ruta: str | Path) -> Path:
    ruta = Path(ruta)
    e = r.examen
    est = r.por_estudiante.copy()
    est["aprobado"] = est["aprobado"].map({True: "Sí", False: "No"})
    est = est.rename(columns={"puesto": "Puesto", "forma": "Forma", "documento": "Documento", "nombre": "Nombre",
                              "curso": "Curso", "aciertos": "Aciertos",
                              "porcentaje": "% total", "nota": "Nota",
                              "puntaje": "Puntaje tipo Saber (0–500)",
                              "aprobado": "Aprueba", "archivo": "Hoja"})

    materias = r.por_materia.rename(columns={
        "materia": "Materia", "preguntas": "Preguntas", "promedio": "Promedio %",
        "maximo": "Máximo %", "minimo": "Mínimo %", "nota_promedio": "Nota promedio",
        "aprueban": "Estudiantes que aprueban"})

    preguntas = r.por_pregunta.drop(columns=["etiqueta"])
    if not e.letras_formas:
        preguntas = preguntas.drop(columns=["forma"])
    preguntas = preguntas.rename(columns={"forma": "Forma", "respondieron": "Estudiantes",
        "discriminacion": "Discriminación", "calidad": "¿Discrimina?",
        "sospecha_clave": "¿Clave errada? (opción sugerida)",
        "pregunta": "Pregunta", "materia": "Materia", "clave": "Clave", "acierto": "% acierto",
        "dificultad": "Dificultad", "en_blanco": "En blanco", "doble_marca": "Doble marca",
        "distractor": "Opción errada más elegida", "anulada": "Anulada"})
    preguntas["Anulada"] = preguntas["Anulada"].map(
        {"bien": "Sí: buena para todos", "excluir": "Sí: no cuenta"}).fillna("")

    respuestas = r.respuestas.copy()
    numeros = e.numeros_impresos()
    respuestas.columns = [f"P{numeros[c]}" for c in respuestas.columns]
    respuestas.insert(0, "Nombre", [dict(zip(r.por_estudiante["documento"],
                                             r.por_estudiante["nombre"])).get(d, "")
                                    for d in respuestas.index])
    if e.letras_formas:
        respuestas.insert(1, "Forma", r.formas.reindex(respuestas.index))
    respuestas = respuestas.reset_index().rename(columns={"documento": "Documento"})
    primera_respuesta = 4 if e.letras_formas else 3   # columna donde empiezan P1, P2...

    with pd.ExcelWriter(ruta, engine="openpyxl") as escritor:
        # --- Resumen ---
        n = len(r.por_estudiante)
        resumen = pd.DataFrame([
            ("Examen", e.nombre), ("Grado", e.grado),
            ("Estudiantes calificados", n),
            ("Promedio general (%)", round(r.por_estudiante["porcentaje"].mean(), 1)),
            ("Nota promedio", r.escala.nota(r.por_estudiante["porcentaje"].mean() / 100)),
            ("Aprueban", f"{int(r.por_estudiante['aprobado'].sum())} de {n}"),
            *([("Puntaje tipo Saber promedio (0–500)",
                round(r.por_estudiante["puntaje"].mean())),
               ("Pesos del puntaje", ", ".join(
                   f"{m.codigo} ×{pesos_saber(e, r.escala)[m.codigo]:g}" for m in e.materias)),
               ("Nota sobre el puntaje", "Aproximación: cada materia = % de aciertos "
                                         "(el ICFES usa el modelo Rasch).")]
              if r.con_puntaje else []),
            ("Escala de notas", f"{r.escala.minima} a {r.escala.maxima} "
                                f"(aprueba con {r.escala.aprobacion})"),
            ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ], columns=["Dato", "Valor"])
        resumen.to_excel(escritor, sheet_name="Resumen", index=False)
        materias.to_excel(escritor, sheet_name="Resumen", index=False, startrow=len(resumen) + 3)
        est.to_excel(escritor, sheet_name="Estudiantes", index=False)
        materias.to_excel(escritor, sheet_name="Materias", index=False)
        preguntas.to_excel(escritor, sheet_name="Preguntas", index=False)
        respuestas.to_excel(escritor, sheet_name="Respuestas", index=False)
        if r.alertas:
            pd.DataFrame(r.alertas, columns=["Tipo", "Detalle"]).to_excel(
                escritor, sheet_name="Alertas", index=False)

        libro = escritor.book

        # Formato: Resumen
        hoja = libro["Resumen"]
        _formatear_hoja(hoja, {"A": 26, "B": 40})
        for fila in range(2, len(resumen) + 2):
            hoja.cell(row=fila, column=2).alignment = Alignment(horizontal="left")
        fila_materias = len(resumen) + 4
        for celda in hoja[fila_materias]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=VERDE)
        hoja.cell(row=fila_materias - 1, column=1, value="Resultados por materia").font = Font(bold=True, size=12)
        # Gráfico de barras: promedio por materia
        grafico = BarChart()
        grafico.title = "Promedio de aciertos por materia (%)"
        grafico.y_axis.scaling.min, grafico.y_axis.scaling.max = 0, 100
        grafico.legend = None
        grafico.height, grafico.width = 7.5, 16
        datos = Reference(hoja, min_col=3, min_row=fila_materias,
                          max_row=fila_materias + len(materias))
        categorias = Reference(hoja, min_col=1, min_row=fila_materias + 1,
                               max_row=fila_materias + len(materias))
        grafico.add_data(datos, titles_from_data=True)
        grafico.set_categories(categorias)
        grafico.series[0].graphicalProperties.solidFill = VERDE
        hoja.add_chart(grafico, f"A{fila_materias + len(materias) + 3}")

        # Formato: Estudiantes (verde aprueba, rojo no)
        hoja = libro["Estudiantes"]
        _formatear_hoja(hoja)
        col_aprueba = list(est.columns).index("Aprueba") + 1
        for fila in hoja.iter_rows(min_row=2):
            color = VERDE_SUAVE if fila[col_aprueba - 1].value == "Sí" else ROJO_SUAVE
            for celda in fila:
                celda.fill = PatternFill("solid", fgColor=color)

        _formatear_hoja(libro["Materias"])

        # Formato: Preguntas (color según dificultad)
        hoja = libro["Preguntas"]
        _formatear_hoja(hoja)
        col_dif = list(preguntas.columns).index("Dificultad") + 1
        colores = {"Fácil": VERDE_SUAVE, "Media": AMARILLO_SUAVE, "Difícil": ROJO_SUAVE}
        col_sosp = list(preguntas.columns).index("¿Clave errada? (opción sugerida)") + 1
        col_anulada = list(preguntas.columns).index("Anulada") + 1
        for fila in hoja.iter_rows(min_row=2):
            fila[col_dif - 1].fill = PatternFill("solid", fgColor=colores[fila[col_dif - 1].value])
            if fila[col_sosp - 1].value:      # posible error en la clave: toda la fila resaltada
                for celda in fila:
                    celda.fill = PatternFill("solid", fgColor="FFE08A")
            if fila[col_anulada - 1].value:   # pregunta anulada: toda la fila en gris
                for celda in fila:
                    celda.fill = PatternFill("solid", fgColor=GRIS)

        # Formato: Respuestas (verde acierto, rojo error, gris en blanco)
        hoja = libro["Respuestas"]
        _formatear_hoja(hoja, {"A": 14, "B": 30})
        aciertos = r.aciertos.reindex(respuestas["Documento"])
        for i, fila in enumerate(hoja.iter_rows(min_row=2, min_col=primera_respuesta), start=0):
            for j, celda in enumerate(fila):
                valor = aciertos.iat[i, j]
                if pd.isna(valor):                       # pregunta sin clave: no se califica
                    continue
                if celda.value in (None, ""):
                    celda.fill = PatternFill("solid", fgColor=GRIS)
                else:
                    celda.fill = PatternFill("solid", fgColor=VERDE_SUAVE if valor else ROJO_SUAVE)
                celda.alignment = Alignment(horizontal="center")
        for j in range(primera_respuesta, hoja.max_column + 1):
            hoja.column_dimensions[get_column_letter(j)].width = 5

        if r.alertas:
            _formatear_hoja(libro["Alertas"], {"A": 22, "B": 90})
    return ruta


def exportar_seguimiento(s, ruta: str | Path) -> Path:
    """
    Excel del seguimiento entre exámenes:
        Curso        → promedios por examen + gráfico de líneas
        Estudiantes  → % de cada estudiante en cada examen y su cambio
    """
    from openpyxl.chart import LineChart

    ruta = Path(ruta)
    curso = s.curso.reset_index().rename(columns={"index": "Examen"})
    est = s.estudiantes.reset_index().rename(columns={
        "documento": "Documento", "nombre": "Nombre", "presentados": "Exámenes presentados",
        "cambio": "Cambio (último − primero)"})
    with pd.ExcelWriter(ruta, engine="openpyxl") as escritor:
        curso.to_excel(escritor, sheet_name="Curso", index=False)
        est.to_excel(escritor, sheet_name="Estudiantes", index=False)
        libro = escritor.book

        hoja = libro["Curso"]
        _formatear_hoja(hoja, {"A": 34})
        # Gráfico: promedio general y por materia en cada examen
        materias = [c for c in curso.columns
                    if c not in ("Examen", "Estudiantes", "Nota promedio", "Aprueban %")]
        grafico = LineChart()
        grafico.title = "Promedio del curso por examen (%)"
        grafico.y_axis.scaling.min, grafico.y_axis.scaling.max = 0, 100
        grafico.height, grafico.width = 9, 20
        for col in materias:
            j = list(curso.columns).index(col) + 1
            grafico.add_data(Reference(hoja, min_col=j, min_row=1, max_row=len(curso) + 1),
                             titles_from_data=True)
        grafico.set_categories(Reference(hoja, min_col=1, min_row=2, max_row=len(curso) + 1))
        hoja.add_chart(grafico, f"A{len(curso) + 4}")

        hoja = libro["Estudiantes"]
        _formatear_hoja(hoja, {"A": 14, "B": 32})
        col_cambio = list(est.columns).index("Cambio (último − primero)") + 1
        for fila in hoja.iter_rows(min_row=2):
            valor = fila[col_cambio - 1].value
            if isinstance(valor, (int, float)) and not pd.isna(valor):
                if valor >= 5:
                    fila[col_cambio - 1].fill = PatternFill("solid", fgColor=VERDE_SUAVE)
                elif valor <= -5:
                    fila[col_cambio - 1].fill = PatternFill("solid", fgColor=ROJO_SUAVE)
    return ruta
