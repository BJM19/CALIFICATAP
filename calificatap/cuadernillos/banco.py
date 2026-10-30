"""
banco.py — El banco de preguntas en Excel.

El docente escribe las preguntas UNA sola vez en un Excel que crea el programa
(banco_preguntas.xlsx). Tiene dos hojas:

    Preguntas
    ┌─────────┬─────┬──────────┬─────────────────┬────────┬─────┬─────┬─────┬─────┬──────────┬──────┐
    │ Materia │ N.º │ Contexto │ Enunciado       │ Imagen │  A  │  B  │  C  │  D  │ Correcta │ Fija │
    ├─────────┼─────┼──────────┼─────────────────┼────────┼─────┼─────┼─────┼─────┼──────────┼──────┤
    │ MAT     │  1  │          │ ¿Cuánto es ...? │        │ 12  │ 15  │ 18  │ 21  │    B     │      │
    │ LEC     │ 26  │ T1       │ Según el texto… │        │ ... │ ... │ ... │ ... │    D     │      │
    └─────────┴─────┴──────────┴─────────────────┴────────┴─────┴─────┴─────┴─────┴──────────┴──────┘

    Contextos (textos, tablas o gráficas que comparten varias preguntas)
    ┌────┬──────────────────────────────┬────────┐
    │ ID │ Texto                        │ Imagen │
    │ T1 │ «Había una vez...»           │        │
    └────┴──────────────────────────────┴────────┘

Las filas de "Preguntas" ya vienen creadas con la materia y el número (según
el diseño de la hoja de respuestas): solo hay que llenar el texto.

IMÁGENES: hay dos formas de poner una imagen (gráfica, mapa, fórmula...):
    1. Pegarla dentro del Excel, encima de la celda "Imagen" de esa fila.
    2. Escribir el nombre del archivo (ej. "grafica1.png") y guardar el archivo
       en la carpeta "imagenes", junto al Excel.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..examenes import Examen

NOMBRE_BANCO = "banco_preguntas.xlsx"
CARPETA_IMAGENES = "imagenes"
LETRAS = "ABCDEFGH"
VERDE = "1F7A4D"
VERDE_CLARO = "E3F1E9"


@dataclass
class Contexto:
    id: str
    texto: str = ""
    imagen: bytes | None = None          # la imagen ya leída (de un archivo o pegada en Excel)


@dataclass
class Pregunta:
    etiqueta: str                        # MAT_01 (posición en el banco = forma A)
    materia: str                         # código: MAT
    enunciado: str = ""
    opciones: list[str] = field(default_factory=list)   # ["12", "15", "18", "21"]
    correcta: str = ""                   # "B"
    contexto: str = ""                   # ID del contexto ("" si no tiene)
    imagen: bytes | None = None
    fija: bool = False                   # True = no cambiar el orden de sus opciones
    fila: int = 0                        # fila del Excel (para los mensajes de error)


@dataclass
class Banco:
    preguntas: dict[str, Pregunta]       # {etiqueta: Pregunta}, en el orden del examen
    contextos: dict[str, Contexto]
    problemas: list[str]                 # lo que falta o está mal (vacío = todo bien)

    @property
    def completo(self) -> bool:
        return not self.problemas


# ---------------------------------------------------------------------------
# 1. Crear la plantilla
# ---------------------------------------------------------------------------
def crear_plantilla(examen: Examen, ruta: str | Path) -> Path:
    """Crea el Excel con una fila por pregunta (materia y número ya escritos)."""
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    (ruta.parent / CARPETA_IMAGENES).mkdir(exist_ok=True)
    max_opciones = max(m.opciones for m in examen.materias)
    numeros = examen.numeros_impresos()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Preguntas"
    titulos = (["Materia", "N.º", "Contexto", "Enunciado", "Imagen"]
               + list(LETRAS[:max_opciones]) + ["Correcta", "Fija"])
    hoja.append(titulos)
    for m in examen.materias:
        for e in m.etiquetas():
            hoja.append([m.codigo, numeros[e]] + [""] * (len(titulos) - 2))

    # Formato: encabezado verde, columnas anchas, texto que salta de línea
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    anchos = {"A": 9, "B": 6, "C": 10, "D": 60, "E": 16}
    for i, titulo in enumerate(titulos, start=1):
        letra = hoja.cell(row=1, column=i).column_letter
        hoja.column_dimensions[letra].width = anchos.get(letra, 22 if titulo in LETRAS else 10)
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
        fila[0].fill = fila[1].fill = PatternFill("solid", fgColor=VERDE_CLARO)
    hoja.freeze_panes = "D2"
    # Lista desplegable en "Correcta" y "Fija"
    col_correcta = hoja.cell(row=1, column=len(titulos) - 1).column_letter
    col_fija = hoja.cell(row=1, column=len(titulos)).column_letter
    ultima = hoja.max_row
    validacion = DataValidation(type="list", formula1=f'"{",".join(LETRAS[:max_opciones])}"',
                                allow_blank=True)
    validacion.add(f"{col_correcta}2:{col_correcta}{ultima}")
    fija = DataValidation(type="list", formula1='"sí"', allow_blank=True)
    fija.add(f"{col_fija}2:{col_fija}{ultima}")
    hoja.add_data_validation(validacion)
    hoja.add_data_validation(fija)

    contextos = libro.create_sheet("Contextos")
    contextos.append(["ID", "Texto", "Imagen"])
    for celda in contextos[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE)
    contextos.column_dimensions["A"].width = 8
    contextos.column_dimensions["B"].width = 90
    contextos.column_dimensions["C"].width = 18

    ayuda = libro.create_sheet("Instrucciones")
    ayuda.column_dimensions["A"].width = 110
    for linea in [
        "CÓMO LLENAR EL BANCO DE PREGUNTAS",
        "",
        "Hoja «Preguntas»: una fila por pregunta. La materia y el número ya están escritos "
        "(no los cambie): son los de la hoja de respuestas en la FORMA A.",
        "• Enunciado: el texto de la pregunta. Puede tener varios párrafos (Alt+Enter para saltar de línea).",
        "• A, B, C, D...: el texto de cada opción. Deje vacías las que no use.",
        "• Correcta: la letra de la respuesta correcta.",
        "• Fija: escriba «sí» si las opciones NO se deben cambiar de orden "
        "(por ejemplo, si una dice «Todas las anteriores» o van de menor a mayor).",
        "• Imagen: pegue la imagen encima de esta celda, o escriba el nombre del archivo "
        "(por ejemplo grafica1.png) y guárdelo en la carpeta «imagenes» junto a este Excel.",
        "• Contexto: si varias preguntas usan el mismo texto, tabla o gráfica, escriba su ID "
        "(por ejemplo T1) en cada una de ellas. Esas preguntas deben ir SEGUIDAS.",
        "",
        "Hoja «Contextos»: un texto o una imagen que comparten varias preguntas.",
        "• ID: un nombre corto (T1, T2, G1...).  • Texto: el texto completo.  • Imagen: igual que arriba.",
        "",
        "Al generar los cuadernillos, el programa cambia el orden de las preguntas (dentro de cada "
        "materia) y/o de las opciones en las formas B, C... y calcula la clave de cada forma.",
        "Las preguntas que comparten un contexto se mueven juntas, siempre con su texto.",
    ]:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=13, color=VERDE)
    libro.move_sheet("Instrucciones", offset=-2)
    libro.active = 1
    libro.save(ruta)
    return ruta


# ---------------------------------------------------------------------------
# 2. Leer el banco (y revisar que esté completo)
# ---------------------------------------------------------------------------
def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)          # 12.0 → "12"
    return str(valor).strip()


def _imagenes_pegadas(hoja) -> dict[int, bytes]:
    """{fila: bytes} de las imágenes pegadas dentro del Excel (según la fila donde empiezan)."""
    imagenes = {}
    for imagen in getattr(hoja, "_images", []):
        try:
            fila = imagen.anchor._from.row + 1          # openpyxl cuenta desde 0
            datos = imagen._data()
        except Exception:
            continue
        imagenes.setdefault(fila, datos)
    return imagenes


def _leer_imagen(nombre: str, pegada: bytes | None, carpeta: Path,
                 donde: str, problemas: list[str]) -> bytes | None:
    """La imagen de una fila: la pegada en el Excel o el archivo con ese nombre."""
    if pegada:
        return pegada
    if not nombre:
        return None
    ruta = carpeta / CARPETA_IMAGENES / nombre
    if not ruta.exists():
        ruta = carpeta / nombre
    try:
        return ruta.read_bytes()
    except OSError:
        problemas.append(f"{donde}: no se encontró la imagen «{nombre}» en la carpeta «imagenes».")
        return None


def leer_banco(examen: Examen, ruta: str | Path) -> Banco:
    """Lee el Excel y revisa que cada pregunta esté completa y coincida con el examen."""
    ruta = Path(ruta)
    carpeta = ruta.parent
    problemas: list[str] = []
    try:
        libro = load_workbook(ruta)              # (sin read_only: así se leen las imágenes)
    except Exception as error:
        return Banco({}, {}, [f"No se pudo abrir el Excel: {error}. ¿Está abierto en otro programa?"])
    if "Preguntas" not in libro.sheetnames:
        return Banco({}, {}, ["El Excel no tiene la hoja «Preguntas». Use la plantilla del programa."])

    # --- Contextos ---
    contextos: dict[str, Contexto] = {}
    if "Contextos" in libro.sheetnames:
        hoja = libro["Contextos"]
        pegadas = _imagenes_pegadas(hoja)
        for n, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            fila = list(fila) + [None] * 3
            ident, texto, imagen = _texto(fila[0]), _texto(fila[1]), _texto(fila[2])
            if not ident:
                if texto or imagen or n in pegadas:
                    problemas.append(f"Contextos, fila {n}: falta el ID.")
                continue
            contextos[ident.upper()] = Contexto(
                ident.upper(), texto,
                _leer_imagen(imagen, pegadas.get(n), carpeta, f"Contextos, fila {n}", problemas))
            if not texto and not contextos[ident.upper()].imagen:
                problemas.append(f"Contexto {ident}: no tiene texto ni imagen.")

    # --- Preguntas ---
    hoja = libro["Preguntas"]
    titulos = [_texto(c.value) for c in hoja[1]]
    col = {t: i for i, t in enumerate(titulos)}
    faltan = [t for t in ("Materia", "Enunciado", "Correcta") if t not in col]
    if faltan:
        return Banco({}, contextos, [f"A la hoja «Preguntas» le faltan las columnas: {', '.join(faltan)}."])
    pegadas = _imagenes_pegadas(hoja)
    letras_col = [t for t in titulos if len(t) == 1 and t in LETRAS]
    materias = {m.codigo: m for m in examen.materias}
    vistas: dict[str, int] = {}
    preguntas: dict[str, Pregunta] = {}

    for n, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        fila = list(fila) + [None] * len(titulos)
        valor = lambda t: _texto(fila[col[t]]) if t in col else ""   # noqa: E731
        codigo = valor("Materia").upper()
        if not codigo:
            continue
        donde = f"Fila {n} ({codigo} {valor('N.º')})"
        if codigo not in materias:
            problemas.append(f"Fila {n}: la materia «{codigo}» no está en el examen.")
            continue
        m = materias[codigo]
        vistas[codigo] = vistas.get(codigo, 0) + 1
        if vistas[codigo] > m.preguntas:
            problemas.append(f"{donde}: sobra (el examen tiene {m.preguntas} preguntas de {m.nombre}).")
            continue
        etiqueta = m.etiquetas()[vistas[codigo] - 1]
        opciones = [valor(l) for l in letras_col]
        while opciones and not opciones[-1]:
            opciones.pop()                       # las opciones vacías del final no cuentan
        p = Pregunta(etiqueta, codigo, valor("Enunciado"), opciones, valor("Correcta").upper(),
                     valor("Contexto").upper(),
                     _leer_imagen(valor("Imagen"), pegadas.get(n), carpeta, donde, problemas),
                     valor("Fija").lower() in ("sí", "si", "x", "1", "true"), n)
        problemas += [f"{donde}: {x}" for x in problemas_de_pregunta(p, m, contextos)]
        preguntas[etiqueta] = p

    for m in examen.materias:
        if vistas.get(m.codigo, 0) < m.preguntas:
            problemas.append(f"{m.nombre}: faltan {m.preguntas - vistas.get(m.codigo, 0)} "
                             f"filas (el examen tiene {m.preguntas} preguntas).")

    orden = [e for e in examen.todas_las_etiquetas() if e in preguntas]
    ordenadas = {e: preguntas[e] for e in orden}
    problemas += problemas_de_contextos(examen, ordenadas)
    return Banco(ordenadas, contextos, problemas)


def imagen_en_memoria(datos: bytes) -> io.BytesIO:
    return io.BytesIO(datos)


# ---------------------------------------------------------------------------
# 3. Revisar UNA pregunta (lo usan leer_banco y el editor de preguntas)
# ---------------------------------------------------------------------------
def problemas_de_pregunta(p: Pregunta, m, contextos: dict) -> list[str]:
    """Lo que le falta o tiene mal una pregunta (lista vacía = está lista)."""
    problemas = []
    if not p.enunciado and not p.imagen and not p.contexto:
        problemas.append("falta el enunciado.")
    opciones = list(p.opciones)
    while opciones and not opciones[-1]:
        opciones.pop()
    if "" in opciones:
        problemas.append("hay una opción vacía en medio de las demás.")
    if len(opciones) < 2:
        problemas.append("necesita al menos 2 opciones.")
    elif len(opciones) > m.opciones:
        problemas.append(f"tiene {len(opciones)} opciones, pero la hoja de respuestas de "
                         f"{m.nombre} solo tiene {m.opciones}.")
    if not p.correcta:
        problemas.append("falta la respuesta correcta.")
    elif p.correcta not in LETRAS[:max(2, len(opciones))]:
        problemas.append(f"la respuesta correcta «{p.correcta}» no es una de sus opciones.")
    if p.contexto and p.contexto not in contextos:
        problemas.append(f"el texto compartido «{p.contexto}» no existe.")
    return problemas


def problemas_de_contextos(examen: Examen, preguntas: dict[str, Pregunta]) -> list[str]:
    """Las preguntas de un mismo texto compartido deben ir seguidas y en una sola prueba."""
    numeros = examen.numeros_impresos()
    orden = [e for e in examen.todas_las_etiquetas() if e in preguntas]
    problemas, ultimo, materia_de = [], {}, {}
    for i, e in enumerate(orden):
        ctx = preguntas[e].contexto
        if not ctx:
            continue
        materia = e.split("_")[0]
        if ctx in ultimo and (ultimo[ctx] != i - 1 or materia_de[ctx] != materia):
            problemas.append(f"Las preguntas del texto {ctx} no están seguidas "
                             f"(revise la pregunta {numeros[e]}).")
        ultimo[ctx], materia_de[ctx] = i, materia
    return list(dict.fromkeys(problemas))


def esta_vacia(p: Pregunta) -> bool:
    return not (p.enunciado or p.imagen or p.contexto or any(p.opciones) or p.correcta)


def banco_vacio(examen: Examen) -> Banco:
    """Un banco con todas las preguntas del examen, sin texto (para empezar a escribir)."""
    preguntas = {e: Pregunta(e, m.codigo) for m in examen.materias for e in m.etiquetas()}
    return Banco(preguntas, {}, [])


def guardar_banco(examen: Examen, banco: Banco, ruta: str | Path) -> Path:
    """
    Escribe el banco completo en el Excel (lo usa el editor de preguntas).
    Las imágenes quedan PEGADAS dentro del Excel, así el archivo lleva todo.
    """
    from openpyxl.drawing.image import Image as ImagenExcel
    ruta = Path(ruta)
    temporal = ruta.with_name(ruta.stem + "_guardando.xlsx")
    crear_plantilla(examen, temporal)
    libro = load_workbook(temporal)
    hoja = libro["Preguntas"]
    titulos = [_texto(c.value) for c in hoja[1]]
    col = {t: i + 1 for i, t in enumerate(titulos)}

    def pegar(hoja_excel, datos: bytes, celda: str, fila: int) -> None:
        try:
            imagen = ImagenExcel(io.BytesIO(datos))
        except Exception:
            return
        escala = min(1.0, 110 / max(1, imagen.height))      # miniatura en la celda
        imagen.width, imagen.height = imagen.width * escala, imagen.height * escala
        imagen.anchor = celda
        hoja_excel.add_image(imagen)
        hoja_excel.row_dimensions[fila].height = max(hoja_excel.row_dimensions[fila].height or 15,
                                                     imagen.height * 0.78)

    for fila in range(2, hoja.max_row + 1):
        codigo = _texto(hoja.cell(fila, 1).value)
        # la fila n de la materia corresponde a la etiqueta n de esa materia
        orden = sum(1 for f in range(2, fila) if _texto(hoja.cell(f, 1).value) == codigo)
        materia = next(m for m in examen.materias if m.codigo == codigo)
        p = banco.preguntas.get(materia.etiquetas()[orden])
        if p is None:
            continue
        hoja.cell(fila, col["Contexto"]).value = p.contexto or None
        hoja.cell(fila, col["Enunciado"]).value = p.enunciado or None
        for letra, texto in zip(LETRAS, p.opciones):
            if letra in col:
                hoja.cell(fila, col[letra]).value = texto or None
        hoja.cell(fila, col["Correcta"]).value = p.correcta or None
        hoja.cell(fila, col["Fija"]).value = "sí" if p.fija else None
        if p.imagen:
            pegar(hoja, p.imagen, f"{hoja.cell(fila, col['Imagen']).column_letter}{fila}", fila)

    contextos = libro["Contextos"]
    for i, ctx in enumerate(banco.contextos.values(), start=2):
        contextos.cell(i, 1).value = ctx.id
        contextos.cell(i, 2).value = ctx.texto or None
        contextos.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if ctx.imagen:
            pegar(contextos, ctx.imagen, f"C{i}", i)
    libro.save(temporal)
    temporal.replace(ruta)            # se reemplaza al final: si algo falla, el Excel queda intacto
    return ruta
