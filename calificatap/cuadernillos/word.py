"""
word.py — Importar preguntas desde un documento de Word (.docx).

Muchos docentes ya tienen sus preguntas en Word. Si el documento sigue este
formato sencillo, el programa las pasa al banco de preguntas:

    MATEMÁTICAS                          ← nombre de la prueba (cambia de materia)

    1. ¿Cuánto es el 15 % de 2.400?      ← número + punto (o paréntesis): una pregunta
    A. 160                               ← opciones: letra + punto (o paréntesis)
    B. 240
    *C. 360                              ← un asterisco marca la correcta, o bien...
    D. 36
    Respuesta: C                         ← ...una línea «Respuesta:» (o «Clave:»)

    RESPONDA LAS PREGUNTAS 26 A 29 DE ACUERDO CON LA SIGUIENTE INFORMACIÓN
    (el texto, tabla o imagen compartidos)    ← se asigna a las 4 preguntas siguientes

    - Las imágenes pegadas en el enunciado se vuelven la imagen de la pregunta.
    - Una línea «Fija» deja las opciones de esa pregunta sin barajar.
    - También entiende las listas numeradas automáticas de Word (1., 2., ... y a., b., ...).
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from ..examenes import Examen
from .banco import LETRAS, Banco, Contexto, Pregunta, esta_vacia

RE_PREGUNTA = re.compile(r"^(\d{1,3})\s*[.)\-–](?!\d)\s*(.*)$", re.S)   # «12.» pero no «2.400»
RE_OPCION = re.compile(r"^(\*?)\s*([A-Ha-h])\s*[.)]\s*(.*)$", re.S)
RE_RESPUESTA = re.compile(r"^(respuesta|clave|correcta|rta)\s*(correcta)?\s*[:=.\-]?\s*"
                          r"([A-Ha-h])\b", re.I)
RE_FIJA = re.compile(r"^\(?\s*fija\s*\)?\.?$", re.I)
RE_INSTRUCCION = re.compile(r"responda\s+(las\s+)?preguntas?\s+(n\.?º?\s*)?(\d+)"
                            r"(\s*(a|y|al|hasta)\s*(la\s+)?(\d+))?", re.I)

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _normal(texto: str) -> str:
    """'Ciencias Naturales:' → 'ciencias naturales' (sin tildes ni signos)."""
    texto = unicodedata.normalize("NFD", texto.lower())
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return " ".join(re.sub(r"[^a-zñ ]", " ", texto).split())


@dataclass
class ResultadoWord:
    preguntas: dict[str, Pregunta] = field(default_factory=dict)
    contextos: dict[str, Contexto] = field(default_factory=dict)
    avisos: list[str] = field(default_factory=list)

    def resumen(self, examen: Examen) -> str:
        partes = []
        for m in examen.materias:
            n = sum(1 for e in self.preguntas if e.startswith(m.codigo + "_"))
            if n:
                partes.append(f"{m.nombre}: {n}")
        return ", ".join(partes)


@dataclass
class _Leida:
    numero: int | None
    materia: str | None
    enunciado: list[str] = field(default_factory=list)
    imagen: bytes | None = None
    opciones: list[str] = field(default_factory=list)
    correcta: str = ""
    fija: bool = False
    contexto: str = ""


def _imagenes_del_parrafo(parrafo, documento) -> list[bytes]:
    imagenes = []
    for blip in parrafo._p.findall(".//a:blip", NS):
        rid = blip.get(f"{{{NS['r']}}}embed")
        try:
            imagenes.append(documento.part.related_parts[rid].blob)
        except KeyError:
            continue
    return imagenes


def _nivel_de_lista(parrafo) -> int | None:
    """Si el párrafo es parte de una lista numerada automática de Word, su nivel (0, 1...)."""
    ppr = parrafo._p.pPr
    num = ppr.numPr if ppr is not None else None
    if num is None:
        # la numeración puede venir del estilo («Lista con números», «List Number 2»...)
        estilo = parrafo.style
        while estilo is not None and num is None:
            pstyle = estilo.element.pPr
            num = pstyle.numPr if pstyle is not None else None
            estilo = estilo.base_style
        if num is None:
            return None
        nombre = (parrafo.style.name or "").strip()
        if nombre[-1:].isdigit() and num.ilvl is None:
            return int(nombre[-1]) - 1          # «List Number 2» → nivel 1
    ilvl = num.ilvl
    return int(ilvl.val) if ilvl is not None else 0


def leer_word(examen: Examen, ruta: str | Path) -> ResultadoWord:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as error:
        raise RuntimeError("Falta el paquete «python-docx». Ejecute Instalar.bat de nuevo "
                           "(con internet).") from error

    documento = docx.Document(str(ruta))
    resultado = ResultadoWord()
    materias = {m.codigo: m for m in examen.materias}
    nombres = {}
    for m in examen.materias:
        nombres[_normal(m.nombre)] = m.codigo
        nombres[_normal(m.codigo)] = m.codigo

    leidas: list[_Leida] = []
    materia_actual: str | None = None
    actual: _Leida | None = None
    contexto: dict | None = None        # {"id", "texto": [], "imagen", "cuantas"}
    pendientes_ctx: list[tuple[str, int]] = []   # (id, cuántas preguntas faltan por asignarle)
    estado = ""                          # "", "contexto", "pregunta", "opciones"
    sin_ubicar = 0

    def cerrar_contexto():
        nonlocal contexto
        if contexto:
            resultado.contextos[contexto["id"]] = Contexto(
                contexto["id"], "\n\n".join(t for t in contexto["texto"] if t), contexto["imagen"])
            pendientes_ctx.append((contexto["id"], contexto["cuantas"]))
            contexto = None

    def nueva_pregunta(numero: int | None, texto: str) -> _Leida:
        nonlocal actual, estado
        cerrar_contexto()
        actual = _Leida(numero, materia_actual)
        if pendientes_ctx:
            ident, faltan = pendientes_ctx[0]
            actual.contexto = ident
            if faltan <= 1:
                pendientes_ctx.pop(0)
            else:
                pendientes_ctx[0] = (ident, faltan - 1)
        if texto:
            actual.enunciado.append(texto)
        leidas.append(actual)
        estado = "pregunta"
        return actual

    for elemento in documento.element.body.iterchildren():
        etiqueta = elemento.tag.split("}")[-1]
        if etiqueta == "tbl":
            resultado.avisos.append("Hay una tabla en el Word: no se importó. Péguela como "
                                    "imagen en la pregunta o en el texto compartido.")
            Table(elemento, documento)          # (solo para validar que se puede leer)
            continue
        if etiqueta != "p":
            continue
        parrafo = Paragraph(elemento, documento)
        texto = parrafo.text.strip()
        imagenes = _imagenes_del_parrafo(parrafo, documento)
        nivel = _nivel_de_lista(parrafo)
        if not texto and not imagenes:
            continue

        normal = _normal(texto) if texto else ""
        # 1) ¿Nombre de una prueba? (línea corta que coincide con una materia)
        codigo = None
        if texto and len(texto) <= 45 and not RE_PREGUNTA.match(texto) and "?" not in texto:
            sin_prueba = normal.replace("prueba de ", "").replace("prueba ", "")
            for nombre, cod in nombres.items():
                if sin_prueba == nombre or (len(nombre) > 3 and sin_prueba.startswith(nombre + " ")):
                    codigo = cod
        if codigo:
            cerrar_contexto()
            materia_actual, actual, estado = codigo, None, ""
            continue
        # 2) «RESPONDA LAS PREGUNTAS 26 A 29 ...»: empieza un texto compartido
        coincide = RE_INSTRUCCION.match(texto) if texto else None
        if coincide:
            cerrar_contexto()
            desde = int(coincide.group(3))
            hasta = int(coincide.group(7)) if coincide.group(7) else desde
            contexto = {"id": f"T{len(resultado.contextos) + len(pendientes_ctx) + 1}",
                        "texto": [], "imagen": None, "cuantas": max(1, hasta - desde + 1)}
            actual, estado = None, "contexto"
            continue
        # 3) Respuesta / Fija
        coincide = RE_RESPUESTA.match(texto) if texto else None
        if coincide and actual:
            actual.correcta = coincide.group(3).upper()
            continue
        if texto and RE_FIJA.match(texto) and actual:
            actual.fija = True
            continue
        if texto and re.match(r"^(respuesta|clave|correcta)\b", texto, re.I):
            continue                      # «Respuesta:» sin letra: se ignora
        # 4) Opción (escrita «A.» o lista automática de nivel 2)
        coincide = RE_OPCION.match(texto) if texto else None
        if actual and (coincide or (nivel is not None and nivel >= 1)):
            if coincide:
                marca, _, contenido = coincide.groups()
            else:
                marca, contenido = ("*" if texto.startswith("*") else ""), texto.lstrip("* ")
            if marca:
                actual.correcta = LETRAS[len(actual.opciones)]
            actual.opciones.append(contenido.strip())
            if imagenes:
                resultado.avisos.append("Una opción tiene una imagen: las opciones solo pueden "
                                        "tener texto (la imagen no se importó).")
            estado = "opciones"
            continue
        # 5) Pregunta nueva (escrita «12.» o lista automática de nivel 1)
        coincide = RE_PREGUNTA.match(texto) if texto else None
        if coincide or nivel == 0:
            numero = int(coincide.group(1)) if coincide else None
            p = nueva_pregunta(numero, (coincide.group(2) if coincide else texto).strip())
            if imagenes:
                p.imagen = imagenes[0]
            continue
        # 6) Texto o imagen sueltos: según dónde estemos
        if estado == "contexto" and contexto is not None:
            if texto:
                contexto["texto"].append(texto)
            if imagenes and contexto["imagen"] is None:
                contexto["imagen"] = imagenes[0]
        elif actual and estado == "pregunta":
            if texto:
                actual.enunciado.append(texto)
            if imagenes and actual.imagen is None:
                actual.imagen = imagenes[0]
        elif actual and estado == "opciones" and texto and actual.opciones:
            actual.opciones[-1] += " " + texto
        # (títulos, instrucciones generales, etc.: se ignoran)
    cerrar_contexto()

    # --- Ubicar cada pregunta en su lugar del examen ---
    numeros = examen.numeros_impresos()
    por_numero = {}
    for e, n in numeros.items():
        por_numero.setdefault(n, []).append(e)
    usadas: set[str] = set()
    siguiente: dict[str, int] = {}
    for leida in leidas:
        etiqueta = None
        if leida.materia:
            lista = materias[leida.materia].etiquetas()
            if leida.numero is not None:
                etiqueta = next((e for e in lista if numeros[e] == leida.numero), None)
                if etiqueta is None and 1 <= leida.numero <= len(lista):
                    etiqueta = lista[leida.numero - 1]      # numeración desde 1 en cada prueba
            if etiqueta is None or etiqueta in usadas:
                i = siguiente.get(leida.materia, 0)
                libres = [e for e in lista[i:] if e not in usadas]
                etiqueta = libres[0] if libres else None
        elif leida.numero is not None:
            candidatas = [e for e in por_numero.get(leida.numero, []) if e not in usadas]
            etiqueta = candidatas[0] if candidatas else None
        if etiqueta is None:
            sin_ubicar += 1
            continue
        usadas.add(etiqueta)
        codigo = etiqueta.split("_")[0]
        siguiente[codigo] = materias[codigo].etiquetas().index(etiqueta) + 1
        resultado.preguntas[etiqueta] = Pregunta(
            etiqueta, codigo, "\n".join(leida.enunciado), leida.opciones, leida.correcta,
            leida.contexto, leida.imagen, leida.fija)
        if not leida.correcta:
            resultado.avisos.append(f"Pregunta {numeros[etiqueta]}: no se encontró la respuesta "
                                    "correcta (márquela con * o «Respuesta: X»).")
    if sin_ubicar:
        resultado.avisos.append(f"{sin_ubicar} pregunta(s) no se pudieron ubicar: su número no "
                                "existe en el examen o falta el nombre de la prueba antes.")
    # quitar contextos que no quedaron en ninguna pregunta
    usados_ctx = {p.contexto for p in resultado.preguntas.values()}
    resultado.contextos = {k: v for k, v in resultado.contextos.items() if k in usados_ctx}
    resultado.avisos = list(dict.fromkeys(resultado.avisos))     # sin repetidos
    return resultado


def combinar(banco: Banco, leido: ResultadoWord, reemplazar: bool) -> tuple[Banco, int]:
    """
    Pasa lo leído del Word al banco. Con reemplazar=False solo llena las preguntas
    que estaban vacías. Los textos compartidos del Word reciben IDs nuevos si chocan.
    Devuelve (banco, cuántas preguntas se pasaron).
    """
    preguntas = dict(banco.preguntas)
    contextos = dict(banco.contextos)
    renombre = {}
    for ident, ctx in leido.contextos.items():
        nuevo = ident
        k = len(contextos) + 1
        while nuevo in contextos:
            nuevo = f"T{k}"
            k += 1
        renombre[ident] = nuevo
        contextos[nuevo] = Contexto(nuevo, ctx.texto, ctx.imagen)
    pasadas = 0
    for e, p in leido.preguntas.items():
        if e in preguntas and (reemplazar or esta_vacia(preguntas[e])):
            p.contexto = renombre.get(p.contexto, p.contexto)
            preguntas[e] = p
            pasadas += 1
    usados = {p.contexto for p in preguntas.values()}
    contextos = {k: v for k, v in contextos.items() if k in usados}
    return Banco(preguntas, contextos, []), pasadas


def crear_plantilla_word(examen: Examen, ruta: str | Path) -> Path:
    """Un Word de ejemplo con las instrucciones y el formato que entiende el programa."""
    import docx
    from docx.shared import Pt, RGBColor

    ruta = Path(ruta)
    documento = docx.Document()
    estilo = documento.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    titulo = documento.add_heading("Preguntas para el cuadernillo", level=1)
    titulo.runs[0].font.color.rgb = RGBColor(0x1F, 0x7A, 0x4D)
    ayuda = documento.add_paragraph()
    ayuda.add_run("Cómo escribir las preguntas (borre esta parte cuando termine):\n").bold = True
    ayuda.add_run(
        "• Antes de las preguntas de cada prueba, escriba su nombre en una línea sola "
        "(por ejemplo MATEMÁTICAS).\n"
        "• Cada pregunta empieza con su número y un punto: «1. ...».\n"
        "• Las opciones empiezan con su letra y un punto: «A. ...».\n"
        "• Marque la correcta con un asterisco (*C. ...) o escriba debajo «Respuesta: C».\n"
        "• Para un texto que comparten varias preguntas, escriba antes: «RESPONDA LAS "
        "PREGUNTAS 25 A 28 DE ACUERDO CON LA SIGUIENTE INFORMACIÓN» y luego el texto.\n"
        "• Puede pegar imágenes (gráficas, fórmulas) en el enunciado o en el texto compartido.\n"
        "• Escriba «Fija» debajo de una pregunta si sus opciones no se deben cambiar de orden.")
    numeros = examen.numeros_impresos()
    for i, m in enumerate(examen.materias):
        documento.add_paragraph()
        documento.add_paragraph().add_run(m.nombre.upper()).bold = True
        primero = numeros[m.etiquetas()[0]]
        if i == 0:
            documento.add_paragraph(f"{primero}. ¿Cuál es el 15 % de 2.400?")
            for letra, texto in zip("ABCD", ["160", "240", "360", "36"]):
                documento.add_paragraph(f"{'*' if letra == 'C' else ''}{letra}. {texto}")
            documento.add_paragraph(f"{primero + 1}. Si 3x + 5 = 20, ¿cuál es el valor de x?")
            for letra, texto in zip("ABCD", ["5", "15", "25/3", "3"]):
                documento.add_paragraph(f"{letra}. {texto}")
            documento.add_paragraph("Respuesta: A")
        else:
            documento.add_paragraph(f"RESPONDA LAS PREGUNTAS {primero} A {primero + 1} DE "
                                    "ACUERDO CON LA SIGUIENTE INFORMACIÓN")
            documento.add_paragraph("(Escriba aquí el texto, o pegue la tabla o gráfica como imagen.)")
            for k in range(2):
                documento.add_paragraph(f"{primero + k}. (Escriba aquí la pregunta)")
                for letra in LETRAS[:m.opciones]:
                    documento.add_paragraph(f"{letra}. ")
                documento.add_paragraph("Respuesta: ")
    documento.save(str(ruta))
    return ruta
